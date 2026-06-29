from pathlib import Path

from openpyxl import Workbook

from uni_intel.ingestion.parsers.student import StudentCompletionsParser, StudentSectionParser, StudentSummaryParser


def test_student_summary_parser_reads_provider_metrics(tmp_path: Path) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    ws4 = workbook.create_sheet("4")
    ws9 = workbook.create_sheet("9")
    for ws in [ws4, ws9]:
        ws.append([])
        ws.append([])
        ws.append([" ", " ", "Commencing Students", None, None, "All Students"])
        ws.append(["State", "Institution", 2023, 2024, "%", 2023, 2024, "%"])
        ws.append(["New South Wales", "The University of Sydney", 10, 12, 0.2, 50, 60, 0.2])
        ws.append(["New South Wales", "Total New South Wales", 10, 12, 0.2, 50, 60, 0.2])
    path = tmp_path / "student_summary.xlsx"
    workbook.save(path)

    rows = StudentSummaryParser().parse(path)

    assert len(rows) == 8
    assert rows[0].source_provider_name == "The University of Sydney"
    assert {row.metric_id for row in rows} == {
        "student_commencing_enrolments",
        "student_total_enrolments",
        "student_commencing_load_eftsl",
        "student_total_load_eftsl",
    }


def test_student_completions_parser_reads_time_series(tmp_path: Path) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    ws = workbook.create_sheet("14.4")
    ws.append([])
    ws.append([])
    ws.append(["State", "Institution", 2023, 2024])
    ws.append(["Victoria", "Monash University", 100, 120])
    path = tmp_path / "completions.xlsx"
    workbook.save(path)

    rows = StudentCompletionsParser().parse(path)

    assert len(rows) == 2
    assert rows[1].metric_id == "student_award_course_completions"
    assert rows[1].reporting_year == 2024
    assert rows[1].numeric_value == 120


def test_student_section_parser_reads_provider_total_and_skips_summary_rows(tmp_path: Path) -> None:
    source = _section_workbook(
        tmp_path,
        sheet_name="1.5",
        title="Table 1.5: Commencing Students by State, Higher Education Institution and Broad Level of Course, 2024",
        total_column_name="Total",
    )

    rows = StudentSectionParser(
        year=2024,
        section=1,
        sheet_name="1.5",
        metric_id="student_commencing_enrolments",
        population="Commencing Students",
        total_column_name="Total",
    ).parse(source)

    assert len(rows) == 1
    assert rows[0].source_provider_name == "The University of Sydney"
    assert rows[0].metric_id == "student_commencing_enrolments"
    assert rows[0].numeric_value == 28783
    assert rows[0].dimensions["student_population"] == "Commencing Students"


def test_student_section_parser_reads_total_eftsl_column(tmp_path: Path) -> None:
    source = _section_workbook(
        tmp_path,
        sheet_name="3.1",
        title="Table 3.1: Actual Student Load (EFTSL) for Commencing Students by State, Higher Education Institution and Broad Level of Course, 2024",
        total_column_name="Total EFTSL",
    )

    rows = StudentSectionParser(
        year=2024,
        section=3,
        sheet_name="3.1",
        metric_id="student_commencing_load_eftsl",
        population="Commencing Students EFTSL",
        total_column_name="Total EFTSL",
    ).parse(source)

    assert len(rows) == 1
    assert rows[0].metric_id == "student_commencing_load_eftsl"
    assert rows[0].numeric_value == 28783


def test_student_section_parser_reads_all_load_total_column(tmp_path: Path) -> None:
    source = _section_workbook(
        tmp_path,
        sheet_name="4.1",
        title="Table 4.1: Actual Student Load (EFTSL) for All Students by State, Higher Education Institution and Broad Level of Course, 2024",
        total_column_name="Total",
    )

    rows = StudentSectionParser(
        year=2024,
        section=4,
        sheet_name="4.1",
        metric_id="student_total_load_eftsl",
        population="All Students EFTSL",
        total_column_name="Total",
    ).parse(source)

    assert len(rows) == 1
    assert rows[0].metric_id == "student_total_load_eftsl"
    assert rows[0].numeric_value == 28783


def _section_workbook(tmp_path: Path, sheet_name: str, title: str, total_column_name: str) -> Path:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    worksheet = workbook.create_sheet(sheet_name)
    worksheet.append(["< Back to Contents >"])
    worksheet.append([title])
    worksheet.append(
        [
            "State",
            "Institution",
            "Postgraduate by Research",
            "Postgraduate by Coursework",
            "Bachelor",
            "Sub-Bachelor",
            "Enabling Courses",
            "Non-award Courses/ Microcredentials",
            total_column_name,
        ]
    )
    worksheet.append(["New South Wales", "The University of Sydney", 1, 2, 3, 4, 5, 6, 28783])
    worksheet.append(["New South Wales", "Total New South Wales", 1, 2, 3, 4, 5, 6, 99999])
    path = tmp_path / f"{sheet_name}.xlsx"
    workbook.save(path)
    return path
