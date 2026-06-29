from pathlib import Path

from openpyxl import Workbook
import pytest

from uni_intel.ingestion.parsers.finance import FinanceParseError, FinanceParser


def test_finance_parser_reads_headerless_csv(tmp_path: Path) -> None:
    source = tmp_path / "finance.csv"
    source.write_text(
        "\n".join(
            [
                "HAROLD:New_Uni_Financials,Total Institution,2024,2024Statements,"
                'The University of Melbourne,"Royalties, Trademarks and Licenses",12345',
                'HAROLD:New_Uni_Financials,HED,2024,2024Statements,'
                '"University of Technology, Sydney",Total Revenues,"1,234"',
            ]
        ),
        encoding="utf-8",
    )

    rows = FinanceParser().parse(source)

    assert rows[0].row_number == 1
    assert rows[0].reporting_year == 2024
    assert rows[0].source_provider_name == "The University of Melbourne"
    assert rows[0].line_item == "Royalties, Trademarks and Licenses"
    assert rows[0].numeric_value == 12345
    assert rows[1].source_provider_name == "University of Technology, Sydney"
    assert rows[1].numeric_value == 1234


def test_finance_parser_rejects_wrong_column_count(tmp_path: Path) -> None:
    source = tmp_path / "bad.csv"
    source.write_text("too,few,columns\n", encoding="utf-8")

    with pytest.raises(FinanceParseError, match="expected 7"):
        FinanceParser().parse(source)


def test_finance_parser_reads_legacy_xlsx_wide_tables(tmp_path: Path) -> None:
    source = tmp_path / "finance.xlsx"
    workbook = Workbook()
    total = workbook.active
    total.title = "Financial Performance - Total"
    total.cell(2, 4, "Harold:New_Uni_Financials")
    total.cell(3, 4, 2021)
    total.cell(4, 4, "2021Statements")
    total.cell(9, 4, "The University of Sydney")
    total.cell(9, 5, "New South Wales")
    total.cell(9, 6, "The University of Melbourne")
    total.cell(11, 2, "Total Revenues from Continuing Operations")
    total.cell(11, 4, 100)
    total.cell(11, 5, 200)
    total.cell(11, 6, 250)

    dual = workbook.create_sheet("Financial Perf - Dual Sector")
    dual.cell(2, 4, "Harold:New_Uni_Financials")
    dual.cell(3, 4, 2021)
    dual.cell(4, 4, "2021Statements")
    dual.cell(7, 4, "Federation University Australia")
    dual.cell(7, 6, "RMIT University")
    dual.cell(7, 8, "Swinburne University of Technology")
    dual.cell(8, 4, "HED")
    dual.cell(8, 5, "Total Institution")
    dual.cell(8, 6, "HED")
    dual.cell(8, 7, "Total Institution")
    dual.cell(8, 8, "HED")
    dual.cell(8, 9, "Total Institution")
    dual.cell(10, 2, "Academic Employee Expenses")
    dual.cell(10, 4, 300)
    dual.cell(10, 5, 400)
    dual.cell(10, 6, 500)
    dual.cell(10, 7, 600)
    dual.cell(10, 8, 700)
    dual.cell(10, 9, 800)
    workbook.save(source)

    rows = FinanceParser().parse(source)

    assert rows[0].reporting_year == 2021
    assert rows[0].source_provider_name == "The University of Sydney"
    assert rows[0].line_item == "Total Revenues from Continuing Operations"
    assert rows[0].numeric_value == 100
    assert any(
        row.source_provider_name == "Federation University Australia"
        and row.institution_scope == "HED"
        and row.line_item == "Academic Employee Expenses"
        and row.numeric_value == 300
        for row in rows
    )
    assert not any(
        row.source_provider_name == "Federation University Australia"
        and row.institution_scope == "Total Institution"
        and row.line_item == "Academic Employee Expenses"
        for row in rows
    )
