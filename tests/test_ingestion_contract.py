from pathlib import Path
from shutil import copyfile
from zipfile import ZipFile

import duckdb
from openpyxl import Workbook

from uni_intel.db import init_schema
from uni_intel.ingestion import ingest_qilt as ingest_qilt_module
from uni_intel.ingestion import ingest_student as ingest_student_module
from uni_intel.ingestion.common import upsert_metrics
from uni_intel.ingestion.ingest_finance import ingest_finance
from uni_intel.ingestion.metrics import CALCULATED_METRICS, STUDENT_METRICS
from uni_intel.ingestion.parsers.student import StudentSectionParser
from uni_intel.seed import seed_providers


def test_finance_ingestion_is_idempotent(tmp_path: Path) -> None:
    source = tmp_path / "finance.csv"
    source.write_text(
        "\n".join(
            [
                "HAROLD:New_Uni_Financials,Total Institution,2024,2024Statements,"
                "The University of Sydney,Operating Result from Continuing Operations,10",
                "HAROLD:New_Uni_Financials,Total Institution,2024,2024Statements,"
                "The University of Sydney,Total Revenues from Continuing Operations (including Deferred Superannuation),100",
            ]
        ),
        encoding="utf-8",
    )
    db_path = tmp_path / "test.duckdb"

    ingest_finance(2024, db_path, "local://finance-test", raw_path=source)
    ingest_finance(2024, db_path, "local://finance-test", raw_path=source)

    conn = duckdb.connect(str(db_path))
    try:
        assert conn.execute("SELECT COUNT(*) FROM stg_finance_rows").fetchone()[0] == 2
        assert conn.execute("SELECT COUNT(*) FROM facts").fetchone()[0] == 2
    finally:
        conn.close()


def test_finance_ingestion_accepts_configured_historical_year(tmp_path: Path) -> None:
    source = tmp_path / "finance.csv"
    source.write_text(
        "HAROLD:New_Uni_Financials,Total Institution,2023,2023Statements,"
        "The University of Sydney,Academic Employee Expenses,123\n",
        encoding="utf-8",
    )
    db_path = tmp_path / "test.duckdb"

    result = ingest_finance(2023, db_path, "local://finance-test", raw_path=source)

    conn = duckdb.connect(str(db_path))
    try:
        assert result["reporting_year"] == 2023
        assert conn.execute("SELECT COUNT(*) FROM facts WHERE reporting_year = 2023").fetchone()[0] == 1
        assert conn.execute("SELECT dataset_id FROM source_files").fetchone()[0] == "education_finance_2023"
    finally:
        conn.close()


def test_student_section_ingestion_is_idempotent_and_purges_summary_facts(
    tmp_path: Path,
    monkeypatch,
) -> None:
    source = tmp_path / "student_section_2.xlsx"
    _write_student_section_workbook(source)
    db_path = tmp_path / "test.duckdb"
    _insert_old_student_summary_fact(db_path)

    monkeypatch.setattr(ingest_student_module, "download_file", lambda *args, **kwargs: None)
    monkeypatch.setattr(
        ingest_student_module,
        "build_student_sources",
        lambda: [
            (
                "student_2024_section_2_all_students",
                "local://student-section-2",
                source,
                StudentSectionParser(
                    year=2024,
                    section=2,
                    sheet_name="2.5",
                    metric_id="student_total_enrolments",
                    population="All Students",
                    total_column_name="Total",
                ),
                "2024 Section 2 - All students",
            )
        ],
    )

    ingest_student_module.ingest_student(2024, db_path)
    ingest_student_module.ingest_student(2024, db_path)

    conn = duckdb.connect(str(db_path))
    try:
        assert conn.execute("SELECT COUNT(*) FROM stg_student_rows").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM facts").fetchone()[0] == 1
        assert (
            conn.execute("SELECT COUNT(*) FROM facts WHERE source_file_id = 'old_student_summary'").fetchone()[0] == 0
        )
        assert (
            conn.execute(
                """
            SELECT COUNT(*)
            FROM facts
            WHERE metric_id = 'student_total_enrolments'
              AND source_file_id IS NOT NULL
              AND source_row_number IS NOT NULL
              AND source_line_item IS NOT NULL
              AND dimensions_json IS NOT NULL
            """
            ).fetchone()[0]
            == 1
        )
    finally:
        conn.close()


def test_qilt_ingestion_loads_configured_history_idempotently(
    tmp_path: Path,
    monkeypatch,
) -> None:
    source = tmp_path / "qilt_source.zip"
    workbook_path = tmp_path / "national.xlsx"
    _write_qilt_history_workbook(workbook_path)
    with ZipFile(source, "w") as archive:
        archive.write(workbook_path, "2021 SES National Report Tables.xlsx")

    db_path = tmp_path / "test.duckdb"
    monkeypatch.setattr(ingest_qilt_module, "RAW_DIR", tmp_path / "raw")
    monkeypatch.setattr(ingest_qilt_module, "QILT_SES_URLS", {2021: "local://qilt-2021.zip"})
    monkeypatch.setattr(ingest_qilt_module, "download_file", _copy_qilt_source(source))

    ingest_qilt_module.ingest_qilt(db_path)
    ingest_qilt_module.ingest_qilt(db_path)

    conn = duckdb.connect(str(db_path))
    try:
        assert conn.execute("SELECT COUNT(*) FROM stg_qilt_rows").fetchone()[0] == 12
        assert conn.execute("SELECT COUNT(*) FROM facts").fetchone()[0] == 12
        assert conn.execute("SELECT DISTINCT reporting_year FROM facts").fetchone()[0] == 2021
        assert (
            conn.execute(
                """
            SELECT COUNT(*)
            FROM facts
            WHERE metric_id LIKE 'qilt%'
              AND source_file_id IS NOT NULL
              AND source_row_number IS NOT NULL
              AND source_line_item IS NOT NULL
              AND dimensions_json IS NOT NULL
            """
            ).fetchone()[0]
            == 12
        )
    finally:
        conn.close()


def test_calculated_metrics_store_methods() -> None:
    assert CALCULATED_METRICS
    for metric in CALCULATED_METRICS:
        assert metric["is_calculated"] is True
        assert metric["calculation_method"]


def _write_student_section_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "2.5"
    worksheet.append([])
    worksheet.append(["Table 2.5: All Students by State, Higher Education Institution and Broad Level of Course, 2024"])
    worksheet.append(["State", "Institution", "Bachelor", "Total"])
    worksheet.append(["NSW", "The University of Sydney", 10, 100])
    worksheet.append(["NSW", "Total NSW", 10, 100])
    workbook.save(path)


def _write_qilt_history_workbook(path: Path) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    headers = [
        "",
        "",
        "Skills Development",
        "Learner Engagement",
        "Teaching Quality",
        "Student Support",
        "Learning Resources",
        "Quality of entire educational experience",
    ]
    values = [
        "",
        "The University of Sydney",
        "80.1 (79.0, 81.2)",
        "60.2 (59.0, 61.4)",
        "78.3 (77.1, 79.5)",
        "75.4 (74.1, 76.7)",
        "82.5 (81.2, 83.8)",
        "70.6 (69.4, 71.8)",
    ]
    for sheet_name in ["FOCUS_UG_UNI_1Y_INST_CI", "FOCUS_PGC_UNI_1Y_INST_CI"]:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["Student experience"])
        sheet.append(["Back to INDEX"])
        sheet.append(["", "", "", "", "", "", "", ""])
        sheet.append(headers)
        sheet.append(values)
    workbook.save(path)


def _copy_qilt_source(source: Path):
    def copy_source(_url: str, destination: Path, force: bool = False) -> None:
        destination.parent.mkdir(parents=True, exist_ok=True)
        copyfile(source, destination)

    return copy_source


def _insert_old_student_summary_fact(db_path: Path) -> None:
    conn = duckdb.connect(str(db_path))
    try:
        init_schema(conn)
        seed_providers(conn)
        upsert_metrics(conn, STUDENT_METRICS)
        conn.execute(
            """
            INSERT INTO source_files (
                source_file_id, dataset_id, source_name, source_url,
                local_path, file_format, reporting_year, downloaded_at,
                checksum_sha256, row_count
            )
            VALUES (
                'old_student_summary', 'education_student_2024',
                'Old summary', 'local://old-summary', '/tmp/old.xlsx',
                'xlsx', 2024, CURRENT_TIMESTAMP, 'old-checksum', 1
            )
            """
        )
        conn.execute(
            """
            INSERT INTO facts (
                fact_id, provider_id, metric_id, source_file_id, reporting_year,
                dimension_scope, value, unit, source_row_number,
                source_provider_name, source_line_item, dimensions_json
            )
            VALUES (
                'old_summary_fact', 'university_of_sydney',
                'student_total_enrolments', 'old_student_summary', 2024,
                'Student', 99, 'students', 1, 'The University of Sydney',
                'Total', '{}'
            )
            """
        )
    finally:
        conn.close()
