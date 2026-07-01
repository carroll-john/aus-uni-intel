from __future__ import annotations

import csv
from collections import defaultdict
from dataclasses import dataclass, replace
from pathlib import Path

from openpyxl import load_workbook


class FinanceParseError(ValueError):
    pass


@dataclass(frozen=True)
class FinanceRawRow:
    row_number: int
    source_system: str
    institution_scope: str
    reporting_year: int
    statement_code: str
    source_provider_name: str
    line_item: str
    raw_value: str
    numeric_value: float


class FinanceParser:
    expected_columns = 7

    def parse(self, path: Path | str) -> list[FinanceRawRow]:
        source_path = Path(path)
        if source_path.suffix.lower() in {".xlsx", ".xlsm"}:
            return self._parse_xlsx(source_path)

        return self._parse_csv(source_path)

    def _parse_csv(self, source_path: Path) -> list[FinanceRawRow]:
        rows: list[FinanceRawRow] = []

        with source_path.open(newline="", encoding="utf-8-sig") as handle:
            for row_number, row in enumerate(csv.reader(handle), start=1):
                if len(row) != self.expected_columns:
                    raise FinanceParseError(
                        f"Row {row_number} has {len(row)} columns; expected {self.expected_columns}"
                    )
                rows.append(self._parse_row(row_number, row))

        return rows

    def _parse_xlsx(self, source_path: Path) -> list[FinanceRawRow]:
        workbook = load_workbook(source_path, read_only=False, data_only=True)
        parsed_rows: list[tuple[FinanceRawRow, str]] = []

        for worksheet in workbook.worksheets:
            source_system = str(worksheet.cell(2, 4).value or "").strip()
            reporting_year = parse_year(worksheet.cell(3, 4).value, worksheet.title)
            statement_code = str(worksheet.cell(4, 4).value or f"{reporting_year}Statements").strip()
            provider_row = find_provider_row(worksheet)
            if provider_row is None:
                continue

            is_dual_sector_sheet = "dual" in worksheet.title.lower()
            scope_row = provider_row + 1 if is_dual_sector_sheet else None
            column_headers = build_column_headers(
                worksheet,
                provider_row=provider_row,
                scope_row=scope_row,
                default_scope="Total Institution",
            )
            data_start_row = provider_row + (3 if is_dual_sector_sheet else 2)

            for row_number in range(data_start_row, worksheet.max_row + 1):
                line_item = str(
                    worksheet.cell(row_number, 2).value or worksheet.cell(row_number, 1).value or ""
                ).strip()
                if not line_item:
                    continue

                for column, provider_name, scope in column_headers:
                    if is_dual_sector_sheet and scope == "Total Institution":
                        continue
                    raw_value = worksheet.cell(row_number, column).value
                    if raw_value is None or raw_value == "":
                        continue
                    try:
                        numeric_value = parse_numeric_value(raw_value, row_number)
                    except FinanceParseError:
                        continue
                    parsed_rows.append(
                        (
                            FinanceRawRow(
                                row_number=row_number,
                                source_system=source_system,
                                institution_scope=scope,
                                reporting_year=reporting_year,
                                statement_code=statement_code,
                                source_provider_name=provider_name,
                                line_item=line_item,
                                raw_value=str(raw_value),
                                numeric_value=numeric_value,
                            ),
                            worksheet.title,
                        )
                    )

        return disambiguate_xlsx_line_items(parsed_rows)

    def _parse_row(self, row_number: int, row: list[str]) -> FinanceRawRow:
        source_system, institution_scope, year, statement_code, provider, line_item, value = [
            col.strip() for col in row
        ]

        try:
            reporting_year = int(year)
        except ValueError as exc:
            raise FinanceParseError(f"Row {row_number} has invalid year: {year}") from exc

        return FinanceRawRow(
            row_number=row_number,
            source_system=source_system,
            institution_scope=institution_scope,
            reporting_year=reporting_year,
            statement_code=statement_code,
            source_provider_name=provider,
            line_item=line_item,
            raw_value=value,
            numeric_value=parse_numeric_value(value, row_number),
        )


def parse_numeric_value(value: str, row_number: int | None = None) -> float:
    if isinstance(value, int | float):
        return float(value)

    cleaned = str(value).strip().replace(",", "")
    if cleaned in {"", "-", "na", "n/a", "NA", "N/A"}:
        message = "Finance values must be numeric"
        if row_number is not None:
            message = f"Row {row_number}: {message}"
        raise FinanceParseError(message)
    try:
        return float(cleaned)
    except ValueError as exc:
        message = f"Invalid numeric finance value: {value}"
        if row_number is not None:
            message = f"Row {row_number}: {message}"
        raise FinanceParseError(message) from exc


def parse_year(value: object, source_name: str) -> int:
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise FinanceParseError(f"{source_name} has invalid year: {value}") from exc


def find_provider_row(worksheet) -> int | None:
    best_row: int | None = None
    best_score = 0
    for row_number in range(1, min(worksheet.max_row, 15) + 1):
        score = 0
        for column in range(4, worksheet.max_column + 1):
            value = worksheet.cell(row_number, column).value
            if isinstance(value, str) and value.strip() and value.strip() not in {"HED", "Total Institution"}:
                score += 1
        if score > best_score:
            best_row = row_number
            best_score = score

    return best_row if best_score >= 3 else None


def build_column_headers(
    worksheet,
    provider_row: int,
    scope_row: int | None,
    default_scope: str,
) -> list[tuple[int, str, str]]:
    headers: list[tuple[int, str, str]] = []
    last_provider: str | None = None

    for column in range(4, worksheet.max_column + 1):
        provider_value = worksheet.cell(provider_row, column).value
        provider_name = str(provider_value).strip() if provider_value else ""
        if provider_name:
            last_provider = provider_name
        elif scope_row is not None:
            provider_name = last_provider or ""

        if not provider_name:
            continue

        scope = default_scope
        if scope_row is not None:
            scope_value = worksheet.cell(scope_row, column).value
            scope = str(scope_value).strip() if scope_value else ""
            if not scope:
                continue

        headers.append((column, provider_name, scope))

    return headers


def disambiguate_xlsx_line_items(parsed_rows: list[tuple[FinanceRawRow, str]]) -> list[FinanceRawRow]:
    grouped: defaultdict[tuple[str, str, int, str], list[tuple[FinanceRawRow, str]]] = defaultdict(list)
    for row, sheet_name in parsed_rows:
        grouped[
            (
                row.source_provider_name,
                row.line_item,
                row.reporting_year,
                row.institution_scope,
            )
        ].append((row, sheet_name))

    ambiguous_line_items = {key[1] for key, rows in grouped.items() if len(rows) > 1}
    rows: list[FinanceRawRow] = []
    for row, sheet_name in parsed_rows:
        if row.line_item in ambiguous_line_items:
            rows.append(replace(row, line_item=f"{row.line_item} ({sheet_name})"))
        else:
            rows.append(row)

    regrouped: defaultdict[tuple[str, str, int, str], list[FinanceRawRow]] = defaultdict(list)
    for row in rows:
        regrouped[
            (
                row.source_provider_name,
                row.line_item,
                row.reporting_year,
                row.institution_scope,
            )
        ].append(row)

    still_ambiguous = {
        (
            row.source_provider_name,
            row.line_item,
            row.reporting_year,
            row.institution_scope,
        )
        for row_group in regrouped.values()
        if len(row_group) > 1
        for row in row_group
    }
    return [
        replace(row, line_item=f"{row.line_item} row {row.row_number}")
        if (row.source_provider_name, row.line_item, row.reporting_year, row.institution_scope) in still_ambiguous
        else row
        for row in rows
    ]
