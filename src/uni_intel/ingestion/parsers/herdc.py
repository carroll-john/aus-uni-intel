from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from uni_intel.ingestion.parsers._numeric import parse_optional_numeric


class HerdcParseError(ValueError):
    pass


@dataclass(frozen=True)
class HerdcRawRow:
    row_number: int
    hep_code: str
    source_provider_name: str
    metric_id: str
    reporting_year: int
    raw_value: str
    numeric_value: float
    dimensions: dict[str, object]


class HerdcResearchIncomeParser:
    def parse(self, path: Path | str) -> list[HerdcRawRow]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet_name = "1. Summary by Category"
        if sheet_name not in workbook.sheetnames:
            raise HerdcParseError(f"Worksheet {sheet_name} not found")
        worksheet = workbook[sheet_name]
        metric_columns = [
            ("herdc_research_income_category_1", 6, "Category 1"),
            ("herdc_research_income_category_2", 7, "Category 2"),
            ("herdc_research_income_category_3", 8, "Category 3"),
            ("herdc_research_income_category_4", 9, "Category 4"),
            ("herdc_research_income_total", 10, "Total"),
        ]

        parsed: list[HerdcRawRow] = []
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=4, values_only=True),
            start=4,
        ):
            hep_code, provider, year, state, cohort = row[:5]
            if not provider or not year:
                continue
            for metric_id, column_index, category in metric_columns:
                raw_value = row[column_index - 1]
                numeric_value = parse_numeric(raw_value)
                if numeric_value is None:
                    continue
                parsed.append(
                    HerdcRawRow(
                        row_number=row_index,
                        hep_code=str(hep_code or ""),
                        source_provider_name=str(provider).strip(),
                        metric_id=metric_id,
                        reporting_year=int(year),
                        raw_value=str(raw_value),
                        numeric_value=numeric_value,
                        dimensions={
                            "state": state,
                            "cohort": cohort,
                            "category": category,
                        },
                    )
                )
        return parsed


def parse_numeric(value: object) -> float | None:
    return parse_optional_numeric(
        value,
        {"", "n/a", "np"},
        lambda bad: HerdcParseError(f"Invalid HERDC numeric value: {bad}"),
    )
