from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

import xlrd
from openpyxl import Workbook, load_workbook

from uni_intel.ingestion.parsers._numeric import parse_optional_numeric


class StudentParseError(ValueError):
    pass


@dataclass(frozen=True)
class StudentRawRow:
    row_number: int
    source_table: str
    sheet_name: str
    source_provider_name: str
    metric_id: str
    reporting_year: int
    raw_value: str
    numeric_value: float
    dimensions: dict[str, object]


@dataclass(frozen=True)
class StudentMetricColumn:
    metric_id: str
    population: str
    column_names: tuple[str, ...]
    level_of_course: str
    fallback_column_names: tuple[tuple[str, ...], ...] = ()


class StudentSummaryParser:
    """Parser for Department 2024 Student Summary Tables workbook."""

    def parse(self, path: Path | str) -> list[StudentRawRow]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows: list[StudentRawRow] = []
        rows.extend(
            self._parse_provider_table(
                workbook,
                sheet_name="4",
                source_table="Table 4: Summary of Student Enrolments by State and Institution - Table A Providers",
                metric_prefix="student",
                metrics=[
                    ("student_commencing_enrolments", 2023, 3, "Commencing Students"),
                    ("student_commencing_enrolments", 2024, 4, "Commencing Students"),
                    ("student_total_enrolments", 2023, 6, "All Students"),
                    ("student_total_enrolments", 2024, 7, "All Students"),
                ],
            )
        )
        rows.extend(
            self._parse_provider_table(
                workbook,
                sheet_name="9",
                source_table="Table 9: Summary of Actual Student Load (EFTSL) by State and Institution - Table A Providers",
                metric_prefix="student",
                metrics=[
                    ("student_commencing_load_eftsl", 2023, 3, "Commencing Students"),
                    ("student_commencing_load_eftsl", 2024, 4, "Commencing Students"),
                    ("student_total_load_eftsl", 2023, 6, "All Students"),
                    ("student_total_load_eftsl", 2024, 7, "All Students"),
                ],
            )
        )
        return rows

    def _parse_provider_table(
        self,
        workbook: Workbook,
        sheet_name: str,
        source_table: str,
        metric_prefix: str,
        metrics: Iterable[tuple[str, int, int, str]],
    ) -> list[StudentRawRow]:
        if sheet_name not in workbook.sheetnames:
            raise StudentParseError(f"Worksheet {sheet_name} not found")

        worksheet = workbook[sheet_name]
        parsed: list[StudentRawRow] = []
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=5, values_only=True),
            start=5,
        ):
            state = row[0]
            provider = row[1]
            if not provider or not isinstance(provider, str):
                continue
            if should_skip_provider_row(provider):
                continue

            for metric_id, year, column_index, population in metrics:
                raw_value = row[column_index - 1]
                numeric_value = parse_numeric(raw_value)
                if numeric_value is None:
                    continue
                parsed.append(
                    StudentRawRow(
                        row_number=row_index,
                        source_table=source_table,
                        sheet_name=sheet_name,
                        source_provider_name=provider,
                        metric_id=metric_id,
                        reporting_year=year,
                        raw_value=str(raw_value),
                        numeric_value=numeric_value,
                        dimensions={
                            "state": state,
                            "provider_table": "Table A Providers",
                            "student_population": population,
                        },
                    )
                )
        return parsed


class StudentSectionParser:
    """Parser for provider-level Department student section workbooks."""

    def __init__(
        self,
        *,
        year: int,
        section: int,
        sheet_name: str,
        metric_id: str,
        population: str,
        total_column_name: str,
        metric_columns: Iterable[StudentMetricColumn] | None = None,
    ) -> None:
        self.year = year
        self.section = section
        self.sheet_name = sheet_name
        self.metric_id = metric_id
        self.population = population
        self.total_column_name = total_column_name
        self.metric_columns = tuple(metric_columns or ())

    def parse(self, path: Path | str) -> list[StudentRawRow]:
        if Path(path).suffix.lower() == ".xls":
            return self._parse_xls(path)

        workbook = load_workbook(path, read_only=True, data_only=True)
        if self.sheet_name not in workbook.sheetnames:
            raise StudentParseError(f"Worksheet {self.sheet_name} not found")

        worksheet = workbook[self.sheet_name]
        title = worksheet.cell(2, 1).value
        header = next(worksheet.iter_rows(min_row=3, max_row=3, values_only=True))
        metric_columns = self._metric_columns(header)
        parsed: list[StudentRawRow] = []

        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=4, values_only=True),
            start=4,
        ):
            state = row[0]
            provider = row[1]
            if not provider or not isinstance(provider, str):
                continue
            if should_skip_provider_row(provider):
                continue

            parsed.extend(
                self._parse_metric_values(
                    row=row,
                    row_number=row_index,
                    state=state,
                    provider=provider,
                    source_table=str(title or f"Section {self.section}"),
                    sheet_name=self.sheet_name,
                    metric_columns=metric_columns,
                )
            )
        return parsed

    def _parse_xls(self, path: Path | str) -> list[StudentRawRow]:
        workbook = xlrd.open_workbook(path)
        sheet_name = find_xls_sheet_name(workbook, f"Table {self.sheet_name}")

        worksheet = workbook.sheet_by_name(sheet_name)
        title_row = find_xls_title_row(worksheet, f"Table {self.sheet_name}")
        header_row = find_xls_header_row(worksheet, self.total_column_name, start=title_row)
        title = worksheet.cell_value(title_row, 0) if title_row is not None else None
        header = tuple(worksheet.row_values(header_row))
        metric_columns = self._metric_columns(header)
        parsed: list[StudentRawRow] = []
        current_state: object = None

        for row_index in range(header_row + 1, worksheet.nrows):
            row = worksheet.row_values(row_index)
            has_any_metric_value = any(
                parse_numeric(row[column_index - 1] if len(row) >= column_index else None) is not None
                for _, _, _, column_indexes in metric_columns
                for column_index in column_indexes
            )

            first_column = row[0] if len(row) > 0 else None
            second_column = row[1] if len(row) > 1 else None
            if not has_any_metric_value and isinstance(first_column, str) and first_column.strip():
                current_state = first_column
                continue

            if isinstance(second_column, str) and second_column.strip():
                state = first_column
                provider = second_column
            else:
                state = current_state
                provider = first_column

            if not provider or not isinstance(provider, str):
                continue
            if should_skip_provider_row(provider):
                continue

            parsed.extend(
                self._parse_metric_values(
                    row=row,
                    row_number=row_index + 1,
                    state=state,
                    provider=provider,
                    source_table=str(title or f"Section {self.section}"),
                    sheet_name=sheet_name,
                    metric_columns=metric_columns,
                )
            )
        return parsed

    def _metric_columns(
        self,
        header: tuple[object, ...],
    ) -> tuple[tuple[str, str, str, tuple[int, ...]], ...]:
        if not self.metric_columns:
            return (
                (
                    self.metric_id,
                    self.population,
                    self.population,
                    (find_header_column(header, self.total_column_name),),
                ),
            )
        return tuple(
            (
                metric.metric_id,
                metric.population,
                metric.level_of_course,
                resolve_metric_column_indexes(header, metric),
            )
            for metric in self.metric_columns
        )

    def _parse_metric_values(
        self,
        *,
        row: tuple[object, ...] | list[object],
        row_number: int,
        state: object,
        provider: str,
        source_table: str,
        sheet_name: str,
        metric_columns: tuple[tuple[str, str, str, tuple[int, ...]], ...],
    ) -> list[StudentRawRow]:
        parsed: list[StudentRawRow] = []
        for metric_id, population, level_of_course, column_indexes in metric_columns:
            values: list[float] = []
            raw_values: list[str] = []
            for column_index in column_indexes:
                raw_value = row[column_index - 1] if len(row) >= column_index else None
                numeric_value = parse_numeric(raw_value)
                if numeric_value is None:
                    continue
                values.append(numeric_value)
                raw_values.append(str(raw_value))
            if not values:
                continue
            parsed.append(
                StudentRawRow(
                    row_number=row_number,
                    source_table=source_table,
                    sheet_name=sheet_name,
                    source_provider_name=provider,
                    metric_id=metric_id,
                    reporting_year=self.year,
                    raw_value=" + ".join(raw_values),
                    numeric_value=sum(values),
                    dimensions={
                        "state": state,
                        "provider_table": "Provider-level section workbook",
                        "student_population": population,
                        "level_of_course": level_of_course,
                    },
                )
            )
        return parsed


class StudentCompletionsParser:
    """Parser for Department Section 14 award course completions workbook."""

    def __init__(
        self,
        *,
        year: int | None = None,
        parse_total_time_series: bool = True,
        parse_level_metrics: bool = False,
    ) -> None:
        self.year = year
        self.parse_total_time_series = parse_total_time_series
        self.parse_level_metrics = parse_level_metrics

    def parse(self, path: Path | str) -> list[StudentRawRow]:
        if Path(path).suffix.lower() == ".xls":
            return self._parse_xls(path)

        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = self._parse_total_time_series_xlsx(workbook) if self.parse_total_time_series else []
        if self.parse_level_metrics:
            rows.extend(self._parse_level_metrics_xlsx(workbook))
        return rows

    def _parse_total_time_series_xlsx(self, workbook: Workbook) -> list[StudentRawRow]:
        sheet_name = "14.4"
        if sheet_name not in workbook.sheetnames:
            raise StudentParseError(f"Worksheet {sheet_name} not found")
        worksheet = workbook[sheet_name]
        header = next(worksheet.iter_rows(min_row=3, max_row=3, values_only=True))
        year_columns: list[tuple[int, int]] = []
        for column_index, value in enumerate(header, start=1):
            if isinstance(value, int) and 2018 <= value <= 2024:
                year_columns.append((column_index, value))

        parsed: list[StudentRawRow] = []
        current_state: object = None
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=4, values_only=True),
            start=4,
        ):
            state = row[0]
            provider = row[1]
            if isinstance(state, str) and state.strip():
                current_state = state
            if not provider and isinstance(state, str) and state.strip():
                provider = state
                state = current_state
            if not provider or not isinstance(provider, str):
                continue
            if should_skip_provider_row(provider):
                continue

            for column_index, year in year_columns:
                raw_value = row[column_index - 1]
                numeric_value = parse_numeric(raw_value)
                if numeric_value is None:
                    continue
                parsed.append(
                    StudentRawRow(
                        row_number=row_index,
                        source_table=(
                            "Table 14.4: Award Course Completions for All Students "
                            "by State and Higher Education Institution"
                        ),
                        sheet_name=sheet_name,
                        source_provider_name=provider,
                        metric_id="student_award_course_completions",
                        reporting_year=year,
                        raw_value=str(raw_value),
                        numeric_value=numeric_value,
                        dimensions={"state": state or current_state, "student_population": "All Students"},
                    )
                )
        return parsed

    def _parse_level_metrics_xlsx(self, workbook: Workbook) -> list[StudentRawRow]:
        sheet_name = "14.8"
        if sheet_name not in workbook.sheetnames:
            raise StudentParseError(f"Worksheet {sheet_name} not found")
        year = self._required_year()
        worksheet = workbook[sheet_name]
        title = worksheet.cell(2, 1).value
        header = next(worksheet.iter_rows(min_row=3, max_row=3, values_only=True))
        metric_columns = completion_metric_columns(header)
        parsed: list[StudentRawRow] = []
        current_state: object = None
        for row_index, row in enumerate(worksheet.iter_rows(min_row=4, values_only=True), start=4):
            state = row[0]
            provider = row[1]
            if isinstance(state, str) and state.strip():
                current_state = state
            if not provider and isinstance(state, str) and state.strip():
                provider = state
                state = current_state
            if not provider or not isinstance(provider, str):
                continue
            if should_skip_provider_row(provider):
                continue
            parsed.extend(
                parse_completion_metric_values(
                    row=row,
                    row_number=row_index,
                    state=state or current_state,
                    provider=provider,
                    year=year,
                    source_table=str(title or "Table 14.8"),
                    sheet_name=sheet_name,
                    metric_columns=metric_columns,
                )
            )
        return parsed

    def _parse_xls(self, path: Path | str) -> list[StudentRawRow]:
        workbook = xlrd.open_workbook(path)
        rows = self._parse_total_time_series_xls(workbook) if self.parse_total_time_series else []
        if self.parse_level_metrics:
            rows.extend(self._parse_level_metrics_xls(workbook))
        return rows

    def _parse_total_time_series_xls(self, workbook: xlrd.book.Book) -> list[StudentRawRow]:
        sheet_name = find_xls_sheet_name(workbook, "Table 14.4")
        worksheet = workbook.sheet_by_name(sheet_name)
        title_row = find_xls_title_row(worksheet, "Table 14.4")
        header_row = find_xls_year_header_row(worksheet, start=title_row, minimum_year=2018)
        title = worksheet.cell_value(title_row, 0) if title_row is not None else "Table 14.4"
        header = tuple(worksheet.row_values(header_row))
        year_columns: list[tuple[int, int]] = []
        for column_index, value in enumerate(header, start=1):
            year = int(value) if isinstance(value, float) and value.is_integer() else value
            if isinstance(year, int) and 2018 <= year <= 2024:
                year_columns.append((column_index, year))

        parsed: list[StudentRawRow] = []
        current_state: object = None
        for row_index in range(header_row + 1, worksheet.nrows):
            row = tuple(worksheet.row_values(row_index))
            first_column = row[0] if len(row) > 0 else None
            second_column = row[1] if len(row) > 1 else None
            has_any_year_value = any(
                parse_numeric(row[column_index - 1] if len(row) >= column_index else None) is not None
                for column_index, _ in year_columns
            )
            if not has_any_year_value and isinstance(first_column, str) and first_column.strip():
                current_state = first_column
                continue
            if isinstance(second_column, str) and second_column.strip():
                state = first_column
                provider = second_column
            else:
                state = current_state
                provider = first_column
            if not provider or not isinstance(provider, str):
                continue
            if should_skip_provider_row(provider):
                continue
            for column_index, year in year_columns:
                raw_value = row[column_index - 1] if len(row) >= column_index else None
                numeric_value = parse_numeric(raw_value)
                if numeric_value is None:
                    continue
                parsed.append(
                    StudentRawRow(
                        row_number=row_index + 1,
                        source_table=str(title or "Table 14.4"),
                        sheet_name=sheet_name,
                        source_provider_name=provider,
                        metric_id="student_award_course_completions",
                        reporting_year=year,
                        raw_value=str(raw_value),
                        numeric_value=numeric_value,
                        dimensions={"state": state, "student_population": "All Students"},
                    )
                )
        return parsed

    def _parse_level_metrics_xls(self, workbook: xlrd.book.Book) -> list[StudentRawRow]:
        year = self._required_year()
        sheet_name = find_xls_sheet_name(workbook, "Table 14.8")
        worksheet = workbook.sheet_by_name(sheet_name)
        title_row = find_xls_title_row(worksheet, "Table 14.8")
        header_row = find_xls_header_row(worksheet, "TOTAL", start=title_row)
        title = worksheet.cell_value(title_row, 0) if title_row is not None else "Table 14.8"
        header = tuple(worksheet.row_values(header_row))
        metric_columns = completion_metric_columns(header)

        parsed: list[StudentRawRow] = []
        current_state: object = None
        for row_index in range(header_row + 1, worksheet.nrows):
            row = tuple(worksheet.row_values(row_index))
            first_column = row[0] if len(row) > 0 else None
            second_column = row[1] if len(row) > 1 else None
            has_any_metric_value = any(
                parse_numeric(row[column_index - 1] if len(row) >= column_index else None) is not None
                for _, _, _, column_indexes in metric_columns
                for column_index in column_indexes
            )
            if not has_any_metric_value and isinstance(first_column, str) and first_column.strip():
                current_state = first_column
                continue
            if isinstance(second_column, str) and second_column.strip():
                state = first_column
                provider = second_column
            else:
                state = current_state
                provider = first_column
            if not provider or not isinstance(provider, str):
                continue
            if should_skip_provider_row(provider):
                continue
            parsed.extend(
                parse_completion_metric_values(
                    row=row,
                    row_number=row_index + 1,
                    state=state,
                    provider=provider,
                    year=year,
                    source_table=str(title or "Table 14.8"),
                    sheet_name=sheet_name,
                    metric_columns=metric_columns,
                )
            )
        return parsed

    def _required_year(self) -> int:
        if self.year is None:
            raise StudentParseError("A reporting year is required for level-of-course completions")
        return self.year


def completion_metric_columns(
    header: tuple[object, ...],
) -> tuple[tuple[str, str, str, tuple[int, ...]], ...]:
    research_broad = find_optional_header_column(header, "Postgraduate by Research")
    coursework_broad = find_optional_header_column(header, "Postgraduate by Coursework")
    if research_broad and coursework_broad:
        return (
            (
                "student_postgraduate_research_completions",
                "Postgraduate by Research",
                "Postgraduate by Research",
                (research_broad,),
            ),
            (
                "student_postgraduate_coursework_completions",
                "Postgraduate by Coursework",
                "Postgraduate by Coursework",
                (coursework_broad,),
            ),
        )

    return (
        (
            "student_postgraduate_research_completions",
            "Postgraduate by Research",
            "Postgraduate by Research",
            tuple(
                find_first_header_column(
                    header,
                    column_names,
                )
                for column_names in (
                    ("Postgraduate by Research", "Higher Doctorate", "Doctorate by Research"),
                    ("Master's by Research", "Masters by Research"),
                )
            ),
        ),
        (
            "student_postgraduate_coursework_completions",
            "Postgraduate by Coursework",
            "Postgraduate by Coursework",
            tuple(
                find_first_header_column(
                    header,
                    column_names,
                )
                for column_names in (
                    ("Postgraduate by Coursework", "Doctorate by Coursework"),
                    ("Master's (Extended)", "Masters (Extended)"),
                    ("Master's by Coursework", "Masters by Coursework"),
                    ("Other Postgraduate",),
                )
            ),
        ),
    )


def parse_completion_metric_values(
    *,
    row: tuple[object, ...],
    row_number: int,
    state: object,
    provider: str,
    year: int,
    source_table: str,
    sheet_name: str,
    metric_columns: tuple[tuple[str, str, str, tuple[int, ...]], ...],
) -> list[StudentRawRow]:
    parsed: list[StudentRawRow] = []
    for metric_id, population, level_of_course, column_indexes in metric_columns:
        values: list[float] = []
        raw_values: list[str] = []
        for column_index in column_indexes:
            raw_value = row[column_index - 1] if len(row) >= column_index else None
            numeric_value = parse_numeric(raw_value)
            if numeric_value is None:
                continue
            values.append(numeric_value)
            raw_values.append(str(raw_value))
        if not values:
            continue
        parsed.append(
            StudentRawRow(
                row_number=row_number,
                source_table=source_table,
                sheet_name=sheet_name,
                source_provider_name=provider,
                metric_id=metric_id,
                reporting_year=year,
                raw_value=" + ".join(raw_values),
                numeric_value=sum(values),
                dimensions={
                    "state": state,
                    "student_population": population,
                    "level_of_course": level_of_course,
                },
            )
        )
    return parsed


def resolve_metric_column_indexes(header: tuple[object, ...], metric: StudentMetricColumn) -> tuple[int, ...]:
    for column_names in (metric.column_names, *metric.fallback_column_names):
        try:
            return tuple(find_header_column(header, column_name) for column_name in column_names)
        except StudentParseError:
            continue
    raise StudentParseError(f"Columns for {metric.metric_id} not found")


def parse_numeric(value: object) -> float | None:
    return parse_optional_numeric(
        value,
        {"", "np", "< 5", "n/a"},
        lambda bad: StudentParseError(f"Invalid numeric value: {bad}"),
    )


def should_skip_provider_row(provider: str) -> bool:
    value = provider.strip()
    normalized = value.lower()
    if not value:
        return True
    if normalized in {"total", "np", "< 5"}:
        return True
    if normalized.startswith(("total ", "% change")):
        return True
    non_public_patterns = (
        "non-university higher education institutions",
        "private universities",
        "avondale university",
        "batchelor institute of indigenous tertiary education",
        "bond university",
        "the university of notre dame australia",
        "torrens university australia",
        "university of divinity",
    )
    return any(pattern in normalized for pattern in non_public_patterns)


def find_header_column(header: tuple[object, ...], expected_name: str) -> int:
    normalized_expected = expected_name.lower()
    for column_index, value in enumerate(header, start=1):
        if str(value or "").strip().lower() == normalized_expected:
            return column_index
    if normalized_expected == "total":
        for column_index, value in enumerate(header, start=1):
            if str(value or "").strip().lower() == "total eftsl":
                return column_index
    raise StudentParseError(f"Column {expected_name} not found")


def find_first_header_column(header: tuple[object, ...], expected_names: tuple[str, ...]) -> int:
    for expected_name in expected_names:
        try:
            return find_header_column(header, expected_name)
        except StudentParseError:
            continue
    raise StudentParseError(f"None of columns {', '.join(expected_names)} found")


def find_optional_header_column(header: tuple[object, ...], expected_name: str) -> int | None:
    try:
        return find_header_column(header, expected_name)
    except StudentParseError:
        return None


def find_xls_sheet_name(workbook: xlrd.book.Book, table_name: str) -> str:
    if table_name.replace("Table ", "") in workbook.sheet_names():
        return table_name.replace("Table ", "")
    for sheet_name in workbook.sheet_names():
        if sheet_name.lower() in {"contents", "content"}:
            continue
        worksheet = workbook.sheet_by_name(sheet_name)
        for row_index in range(min(8, worksheet.nrows)):
            row_text = " ".join(str(value) for value in worksheet.row_values(row_index))
            if table_name in row_text:
                return sheet_name
    raise StudentParseError(f"Worksheet containing {table_name} not found")


def find_xls_title_row(worksheet: xlrd.sheet.Sheet, table_name: str) -> int | None:
    for row_index in range(min(12, worksheet.nrows)):
        row_text = " ".join(str(value) for value in worksheet.row_values(row_index))
        if table_name in row_text:
            return row_index
    return None


def find_xls_header_row(
    worksheet: xlrd.sheet.Sheet,
    total_column_name: str,
    start: int | None,
) -> int:
    start_row = 0 if start is None else start + 1
    for row_index in range(start_row, min(start_row + 8, worksheet.nrows)):
        row = tuple(worksheet.row_values(row_index))
        try:
            find_header_column(row, total_column_name)
            return row_index
        except StudentParseError:
            continue
    raise StudentParseError(f"Header containing {total_column_name} not found")


def find_xls_year_header_row(
    worksheet: xlrd.sheet.Sheet,
    start: int | None,
    minimum_year: int,
) -> int:
    start_row = 0 if start is None else start + 1
    for row_index in range(start_row, min(start_row + 8, worksheet.nrows)):
        for value in worksheet.row_values(row_index):
            year = int(value) if isinstance(value, float) and value.is_integer() else value
            if isinstance(year, int) and year >= minimum_year:
                return row_index
    raise StudentParseError(f"Header containing year {minimum_year} or later not found")
