from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


class QiltParseError(ValueError):
    pass


@dataclass(frozen=True)
class QiltRawRow:
    row_number: int
    source_table: str
    source_provider_name: str
    metric_id: str
    reporting_year: int
    raw_value: str
    numeric_value: float
    ci_lower: float | None
    ci_upper: float | None
    dimensions: dict[str, object]


ODS_NS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}
TABLE = f"{{{ODS_NS['table']}}}"
TEXT = f"{{{ODS_NS['text']}}}"
OFFICE = f"{{{ODS_NS['office']}}}"

QILT_COLUMNS = {
    "Skills Development": "qilt_skills_development_positive_rating",
    "Peer Engagement *": "qilt_peer_engagement_positive_rating",
    "Peer Engagement": "qilt_peer_engagement_positive_rating",
    "Learner Engagement": "qilt_peer_engagement_positive_rating",
    "Teaching Quality and Engagement": "qilt_teaching_quality_engagement_positive_rating",
    "Teaching Quality": "qilt_teaching_quality_engagement_positive_rating",
    "Student Support and Services *": "qilt_student_support_services_positive_rating",
    "Student Support and Services": "qilt_student_support_services_positive_rating",
    "Student Support": "qilt_student_support_services_positive_rating",
    "Learning Resources": "qilt_learning_resources_positive_rating",
    "Quality of entire educational\xa0experience": "qilt_overall_educational_experience_positive_rating",
    "Quality of entire educational experience": "qilt_overall_educational_experience_positive_rating",
}

QILT_PROVIDER_TABLES = {
    "FOCUS_UG_UNI_1Y_INST_CI": "undergraduate",
    "FOCUS_PGC_UNI_1Y_INST_CI": "postgraduate coursework",
}


class QiltSesParser:
    def __init__(self, reporting_year: int = 2024) -> None:
        self.reporting_year = reporting_year

    def parse(self, path: Path | str) -> list[QiltRawRow]:
        tables = _load_tables(Path(path))
        parsed: list[QiltRawRow] = []
        for table_name, course_level in QILT_PROVIDER_TABLES.items():
            table_rows = tables.get(table_name, [])
            if not table_rows:
                raise QiltParseError(f"QILT table {table_name} not found")
            if len(table_rows) < 5:
                raise QiltParseError(f"QILT table {table_name} has no provider rows")

            header_row_index, headers, provider_column = _find_header_row(table_rows)
            metric_columns: list[tuple[int, str, str]] = []
            for index, header in enumerate(headers):
                if header in QILT_COLUMNS:
                    metric_columns.append(
                        (index, QILT_COLUMNS[header], str(header).replace(" *", ""))  # type: ignore[index]
                    )

            for row_number, row in enumerate(table_rows[header_row_index + 1 :], start=header_row_index + 2):
                provider = row[provider_column] if len(row) > provider_column else None
                if not provider or not isinstance(provider, str):
                    continue
                provider = provider.strip()
                for column_index, metric_id, line_item in metric_columns:
                    raw_value = row[column_index] if len(row) > column_index else None
                    parsed_value = parse_qilt_value(raw_value)
                    if parsed_value is None:
                        continue
                    value, ci_lower, ci_upper = parsed_value
                    parsed.append(
                        QiltRawRow(
                            row_number=row_number,
                            source_table=table_name,
                            source_provider_name=provider,
                            metric_id=metric_id,
                            reporting_year=self.reporting_year,
                            raw_value=str(raw_value),
                            numeric_value=value,
                            ci_lower=ci_lower,
                            ci_upper=ci_upper,
                            dimensions={
                                "course_level": course_level,
                                "provider_type": "universities",
                                "confidence_interval": "90%",
                                "ci_lower": ci_lower,
                                "ci_upper": ci_upper,
                                "source_line_item": line_item,
                            },
                        )
                    )
        return parsed


def parse_qilt_value(value: object) -> tuple[float, float | None, float | None] | None:
    if value is None:
        return None
    text = str(value).strip()
    if text.lower() in {"", "n/a", "np", "suppressed"}:
        return None
    match = re.match(r"^([0-9]+(?:\.[0-9]+)?)(?:\s*\(([0-9.]+),\s*([0-9.]+)\))?$", text)
    if not match:
        raise QiltParseError(f"Invalid QILT value: {value}")
    estimate = float(match.group(1))
    lower = float(match.group(2)) if match.group(2) else None
    upper = float(match.group(3)) if match.group(3) else None
    return estimate, lower, upper


def _load_tables(path: Path) -> dict[str, list[list[object | None]]]:
    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            ods_names = _national_names(names, ".ods")
            if not ods_names:
                xlsx_names = _national_names(names, ".xlsx")
                if not xlsx_names:
                    raise QiltParseError("QILT ZIP did not contain a national ODS or XLSX file")
                with archive.open(xlsx_names[0]) as workbook_handle:
                    return _xlsx_tables(workbook_handle.read())
            with archive.open(ods_names[0]) as ods_handle:
                ods_bytes = ods_handle.read()
        return _ods_tables(ods_bytes)

    if path.suffix.lower() == ".xlsx":
        return _xlsx_tables(path.read_bytes())
    return _ods_tables(path.read_bytes())


def _national_names(names: list[str], suffix: str) -> list[str]:
    matching = [name for name in names if name.lower().endswith(suffix)]
    national = [name for name in matching if "national" in name.lower()]
    return national or matching


def _ods_tables(ods_bytes: bytes) -> dict[str, list[list[object | None]]]:
    with zipfile.ZipFile(BytesIO(ods_bytes)) as archive:
        root = ET.fromstring(archive.read("content.xml"))
    return {table.attrib.get(f"{TABLE}name", ""): _table_rows(table) for table in root.findall(f".//{TABLE}table")}


def _xlsx_tables(xlsx_bytes: bytes) -> dict[str, list[list[object | None]]]:
    workbook = load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    return {
        sheet_name: [[cell for cell in row] for row in workbook[sheet_name].iter_rows(values_only=True)]
        for sheet_name in workbook.sheetnames
    }


def _find_header_row(table_rows: list[list[object | None]]) -> tuple[int, list[object | None], int]:
    for row_index, row in enumerate(table_rows):
        metric_columns = [index for index, value in enumerate(row) if value in QILT_COLUMNS]
        if len(metric_columns) >= 3:
            return row_index, row, max(0, min(metric_columns) - 1)
    raise QiltParseError("QILT provider table did not contain expected focus-area headers")


def _load_ods_content(path: Path) -> ET.Element:
    with zipfile.ZipFile(path) as archive:
        return ET.fromstring(archive.read("content.xml"))


def _table_rows(table: ET.Element) -> list[list[object | None]]:
    rows: list[list[object | None]] = []
    for row in table.findall(f"{TABLE}table-row"):
        repeat = int(row.attrib.get(f"{TABLE}number-rows-repeated", "1"))
        values: list[object | None] = []
        for cell in row.findall(f"{TABLE}table-cell"):
            column_repeat = min(
                int(cell.attrib.get(f"{TABLE}number-columns-repeated", "1")),
                64,
            )
            values.extend([_cell_value(cell)] * column_repeat)
        for _ in range(min(repeat, 1)):
            rows.append(values)
    return rows


def _cell_value(cell: ET.Element) -> object | None:
    value = cell.attrib.get(f"{OFFICE}value")
    if value is not None:
        try:
            return float(value)
        except ValueError:
            return value
    string_value = cell.attrib.get(f"{OFFICE}string-value")
    if string_value is not None:
        return string_value
    parts: list[str] = []
    for paragraph in cell.findall(f".//{TEXT}p"):
        text = "".join(paragraph.itertext()).strip()
        if text:
            parts.append(text)
    return " ".join(parts) if parts else None
